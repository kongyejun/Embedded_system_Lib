import os
import re
import time
from dataclasses import dataclass
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from log_system import AppLogger


@dataclass
class BOMGenerateResult:
    success: bool
    input_file: str
    output_file: str
    total_rows: int = 0
    message: str = ""


class BOMGenerator:
    """BOM生成逻辑"""

    def __init__(self, logger: Optional[AppLogger] = None):
        self.logger = logger or AppLogger()

    def generate(self, input_file: str, output_file: Optional[str] = None) -> BOMGenerateResult:
        try:
            if not input_file:
                return BOMGenerateResult(False, "", "", 0, "输入文件路径为空")
            if not os.path.exists(input_file):
                return BOMGenerateResult(False, input_file, "", 0, "输入文件不存在")

            if output_file is None:
                output_file = self._build_output_path(input_file)

            self.logger.info("=" * 60)
            self.logger.info("[BOM生成] 任务开始")
            self.logger.info(f"输入文件：{input_file}")
            self.logger.info(f"输出文件：{output_file}")
            self.logger.info(f"文件大小：{os.path.getsize(input_file)} 字节")

            self.logger.info("步骤1/4：读取源文件...")
            df = self._read_input_file(input_file)
            if df is None:
                self.logger.error("读取失败：无法按Excel/CSV格式解析")
                return BOMGenerateResult(False, input_file, output_file, 0, "文件读取失败，请另存为标准xlsx后重试")

            self.logger.success(f"读取成功：{len(df)} 行，{len(df.columns)} 列")
            self.logger.info(f"列名：{list(df.columns)}")

            self.logger.info("步骤2/4：构建输出工作簿...")
            stats = self._write_bom_workbook(df, input_file, output_file)

            self.logger.info("步骤3/4：保存输出文件...")
            self.logger.success("文件保存成功")

            self.logger.info("步骤4/4：BOM处理结果")
            self.logger.info(f"总处理行数：{len(df)}")
            self.logger.info(f"封装为空行数：{stats['footprint_empty_count']}")
            self.logger.info(f"规则命中统计：{stats['rule_hit']}")
            self.logger.success("【BOM生成】任务完成")
            self.logger.info("=" * 60)

            return BOMGenerateResult(True, input_file, output_file, len(df), "处理成功")

        except Exception as e:
            self.logger.error(f"异常：{str(e)}")
            return BOMGenerateResult(False, input_file, output_file or "", 0, str(e))

    def _build_output_path(self, input_file: str) -> str:
        file_dir = os.path.dirname(input_file)
        file_name = os.path.basename(input_file)
        name_without_ext = os.path.splitext(file_name)[0]
        return os.path.join(file_dir, f"{name_without_ext}_BOM.xlsx")

    def _read_input_file(self, input_file: str) -> Optional[pd.DataFrame]:
        try:
            df = pd.read_excel(input_file, engine="openpyxl")
            self.logger.success("使用 openpyxl 读取Excel成功")
            return df
        except Exception as e:
            self.logger.warn(f"openpyxl读取失败：{e}")

        self.logger.info("尝试CSV编码兜底读取...")
        for encoding in ["utf-8", "gbk", "gb2312", "utf-8-sig"]:
            try:
                df = pd.read_csv(input_file, encoding=encoding)
                self.logger.success(f"CSV读取成功（编码：{encoding}）")
                return df
            except Exception:
                self.logger.debug(f"CSV读取失败（编码：{encoding}）")
                continue
        return None

    def _process_package_and_param(self, package_value, param_value, designator):
        package_str = str(package_value).strip()
        param_str = str(param_value).strip() if pd.notna(param_value) else ""
        designator_str = str(designator).strip().upper() if pd.notna(designator) else ""
        rule = "DEFAULT"

        if package_str.startswith("SMD_RES_CAP"):
            rule = "SMD_RES_CAP"
            if "R" in designator_str:
                return package_str[-4:], "电阻" , param_str, rule
            elif "C" in designator_str:
                return package_str[-4:], "电容" , param_str, rule
            return package_str[-4:]," ",param_str, rule

        elif package_str.startswith("SMD_POL_CAP"):
            rule = "SMD_POL_CAP"
            return package_str[-4:], "钽电容" , param_str, rule

        elif package_str.startswith("SMD_AL_CAP"):
            rule = "SMD_AL_CAP"
            match = re.search(r"SMD_AL_CAP[_-]?(.+)", package_str)
            new_package = match.group(1).replace("-", "x").replace("_", "x") if match else package_str
            return new_package, "铝电解电容" , param_str, rule

        elif package_str.startswith("SMD_RES"):
            rule = "SMD_RES"
            if "W" in package_str.upper():
                return package_str[-4:], "电阻" , param_str, rule

        elif package_str.startswith("LED_SMD"):
            rule = "LED_SMD"
            return package_str[-4:], "",param_str, rule

        return package_str, "",param_str, rule

    def split_order_brand(self,value):
        """
        将类似 'CC0603KRX7R9BB104 国巨' 的字符串拆分为：
        订购号、品牌
        返回:
            code, brand(error_msg)
        """
        if pd.isna(value):
            return "", "!空值"
        text = str(value).strip()
        if not text:
            return "", "!空字符串"
        # 统一处理全角空格、不间断空格、换行、Tab 等
        text = text.replace("\u3000", " ")   # 中文全角空格
        text = text.replace("\xa0", " ")     # 不间断空格
        text = re.sub(r"\s+", " ", text).strip()
        # 正常情况：按第一个空格拆分
        if " " in text:
            parts = text.split(maxsplit=1)
            code = parts[0].strip()
            brand = parts[1].strip()
            return code, brand
        # 非标准情况：没有空格
        # 例如：CC0603KRX7R9BB104国巨
        # 这种情况无法百分百准确判断，只能先整体作为订购号，品牌留空
        return text, "!错误格式请核对"

    def _write_bom_workbook(self, df: pd.DataFrame, input_file: str, output_file: str):
        wb_output = Workbook()
        ws_output = wb_output.active
        ws_output.title = f"BOM_{time.strftime('%Y.%m.%d', time.localtime())}"

        headers = ["序号","品名", "参数规格", "封装", "位号", "数量","推荐品牌","推荐订购号"]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws_output.cell(row=1, column=col_idx, value=header)
            cell.font = Font(name="宋体", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")

        rule_hit = {
            "SMD_RES_CAP": 0,
            "SMD_POL_CAP": 0,
            "SMD_AL_CAP": 0,
            "SMD_RES": 0,
            "LED_SMD": 0,
            "DEFAULT": 0
        }
        footprint_empty_count = 0

        self.logger.info("步骤2.1：写入数据行...")
        for row_idx in range(len(df)):
            out_row = row_idx + 2
            ws_output.cell(row=out_row, column=1, value=row_idx + 1)

            old_param = df.iloc[row_idx, 0] if len(df.columns) > 0 else None
            old_designator = df.iloc[row_idx, 2] if len(df.columns) > 2 else None
            old_footprint = df.iloc[row_idx, 3] if len(df.columns) > 3 else None

            if pd.notna(old_footprint):
                new_footprint,prefix ,param, hit_rule = self._process_package_and_param(
                    old_footprint, old_param, old_designator
                )
                rule_hit[hit_rule] = rule_hit.get(hit_rule, 0) + 1

                ws_output.cell(row=out_row, column=2, value=str(prefix)) # 填写品名
                ws_output.cell(row=out_row, column=3, value=str(param)) # 填写参数规格
                ws_output.cell(row=out_row, column=4, value=new_footprint) # 填写封装

                if str(new_footprint) == str(param):# 若封装与器件名相同则突出显示参数规格
                    ws_output.cell(row=out_row, column=3).font = Font(name="宋体", size=11, color="FF0000")
            else:
                footprint_empty_count += 1# 没有封装则突出显示序号
                ws_output.cell(row=out_row, column=1).font = Font(name="宋体", size=11, color="FF0000")

            if len(df.columns) > 2 and pd.notna(df.iloc[row_idx, 2]):# 填写位号
                ws_output.cell(row=out_row, column=5, value=df.iloc[row_idx, 2])

            if len(df.columns) > 5 and pd.notna(df.iloc[row_idx, 5]):# 填写数量
                ws_output.cell(row=out_row, column=6, value=df.iloc[row_idx, 5])

            if len(df.columns) > 1 and pd.notna(df.iloc[row_idx, 1]):# 填写推荐品牌及订购号
                code, brand = self.split_order_brand(df.iloc[row_idx, 1])
                ws_output.cell(row=out_row, column=7, value= brand)
                if brand[0] == "!":# 突出显示错误
                    ws_output.cell(row=out_row, column=7).font = Font(name="宋体", size=11, color="FF0000")
                ws_output.cell(row=out_row, column=8, value= code)
            
            if (row_idx + 1) % 200 == 0:
                self.logger.info(f"已处理 {row_idx + 1}/{len(df)} 行")

        self.logger.info("步骤2.2：写入末尾板级物料行...")
        last_row = len(df) + 2
        ws_output.cell(row=last_row, column=1, value=len(df) + 1)
        ws_output.cell(row=last_row, column=3, value=os.path.splitext(os.path.basename(input_file))[0])
        ws_output.cell(row=last_row, column=6, value=1)
        for col_idx in range(1, 6):
            ws_output.cell(row=last_row, column=col_idx).fill = PatternFill(
                start_color="FFFF00", end_color="FFFF00", fill_type="solid"
            )

        self.logger.info("步骤2.3：设置列宽...")
        for col_idx in range(1, 7):
            ws_output.column_dimensions[get_column_letter(col_idx)].width = 15

        wb_output.save(output_file)

        return {
            "rule_hit": rule_hit,
            "footprint_empty_count": footprint_empty_count
        }
    

if __name__ == "__main__":
        generator = BOMGenerator(logger=None)
        result = generator.generate("E:\MyFirmCode\BOM快捷生成表软件\VX_XD_T2_TDI_V1.0.xlsx")
