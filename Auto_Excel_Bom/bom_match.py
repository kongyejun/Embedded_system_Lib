import re
from pathlib import Path
from typing import Dict, List, Any, Optional

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Border, Side

from log_system import AppLogger


class BOMMatcher:
    DEFAULT_CONFIG = {
        "pool_target_sheet": "26年正式购买",
        "pool_fallback_sheet_index": 1,
        "output_sheet": "采购表",
        "summary_sheet": "匹配统计",
        "group_headers": ["品牌", "订购号", "最小包装", "物料编码"],
        "header_fill_even": "FFCCCC",
        "header_fill_odd": "CCECFF",
        "apply_all_border": True,
    }

    def __init__(
        self,
        bom_file: Path,
        pool_file: Path,
        config: Optional[Dict[str, Any]] = None,
        logger: Optional[AppLogger] = None
    ):
        self.bom_file = Path(bom_file)
        self.pool_file = Path(pool_file)
        self.cfg = {**self.DEFAULT_CONFIG, **(config or {})}
        self.logger = logger or AppLogger()

        self.bom: Optional[pd.DataFrame] = None
        self.pool: Optional[pd.DataFrame] = None

        self.row_hits: Dict[int, List[List[Any]]] = {}
        self.unique_products: Dict[str, Dict[str, Any]] = {}
        self.stats: Dict[str, Any] = {}

    def run(self) -> Dict[str, Any]:
        self.logger.info("=" * 60)
        self.logger.info("【BOM匹配】任务开始")
        self.logger.info(f"BOM文件：{self.bom_file}")
        self.logger.info(f"汇总池文件：{self.pool_file}")

        self.logger.info("步骤1/3：读取数据...")
        self._load_dataframes()

        self.logger.info("步骤2/3：执行匹配...")
        self._build_row_hits()

        self.logger.info("步骤3/3：写入工作表...")
        self._write_output_sheet()

        self.logger.info("匹配结果汇总")
        self.logger.info(f"总行数：{self.stats.get('total_rows')}")
        self.logger.info(f"命中行数：{self.stats.get('hit_count')}")
        self.logger.info(f"去重后产品数：{self.stats.get('unique_product_count')}")
        self.logger.info(f"重复跳过数：{self.stats.get('repeat_skip_count')}")
        self.logger.info(f"多命中行数：{self.stats.get('multi_count')}")
        self.logger.info(f"未命中行数：{self.stats.get('miss_count')}")
        self.logger.info(f"汇总表重复匹配产品数：{self.stats.get('duplicate_pool_product_count')}")
        self.logger.info(f"最大展开组数：{self.stats.get('max_group_num')}")
        self.logger.success("【BOM匹配】任务完成")
        self.logger.info("=" * 60)
        return self.stats

    def _load_dataframes(self) -> None:
        if not self.bom_file.exists():
            raise FileNotFoundError(f"BOM文件不存在：{self.bom_file}")

        if not self.pool_file.exists():
            raise FileNotFoundError(f"汇总池文件不存在：{self.pool_file}")

        self.bom = pd.read_excel(
            self.bom_file,
            header=0,
            dtype=str,
            sheet_name=0,
            engine="openpyxl"
        )
        self.logger.success(f"BOM读取成功：{len(self.bom)} 行")

        xls = pd.ExcelFile(self.pool_file, engine="openpyxl")
        target_sheet = self.cfg["pool_target_sheet"]
        fallback_idx = self.cfg["pool_fallback_sheet_index"]

        if target_sheet in xls.sheet_names:
            self.pool = pd.read_excel(
                xls,
                sheet_name=target_sheet,
                header=0,
                dtype=str,
                engine="openpyxl"
            )
            self.logger.success(f"汇总池读取成功：sheet={target_sheet}，{len(self.pool)} 行")
        else:
            if len(xls.sheet_names) <= fallback_idx:
                raise ValueError(
                    f"未找到sheet【{target_sheet}】，且回退索引 {fallback_idx} 无效。"
                    f"实际sheet: {xls.sheet_names}"
                )

            fallback_name = xls.sheet_names[fallback_idx]
            self.logger.warn(f"未找到sheet【{target_sheet}】，回退到：{fallback_name}")

            self.pool = pd.read_excel(
                xls,
                sheet_name=fallback_idx,
                header=0,
                dtype=str,
                engine="openpyxl"
            )
            self.logger.success(f"汇总池读取成功：sheet={fallback_name}，{len(self.pool)} 行")

    def _build_row_hits(self) -> None:
        assert self.bom is not None and self.pool is not None

        bom_col_B = self.bom.columns[1]
        bom_col_C = self.bom.columns[2]
        bom_col_D = self.bom.columns[3]

        pool_col_A = self.pool.columns[0]
        pool_col_C = self.pool.columns[2]
        pool_col_D = self.pool.columns[3]
        pool_col_G = self.pool.columns[6]
        pool_col_H = self.pool.columns[7]

        pool_norm_c = self.pool[pool_col_C].map(self._norm)

        row_hits: Dict[int, List[List[Any]]] = {}
        unique_products: Dict[str, Dict[str, Any]] = {}

        hit_count = 0
        multi_count = 0
        miss_count = 0
        repeat_skip_count = 0

        for idx, row in self.bom.iterrows():
            norm_b = self._norm(row[bom_col_B])
            norm_c = self._norm(row[bom_col_C])
            key_b = norm_b + norm_c if norm_b else norm_c
            key_c = self._norm(row[bom_col_D])

            if not key_b or not key_c:
                row_hits[idx] = []
                miss_count += 1
                continue

            cond = (
                pool_norm_c.str.contains(re.escape(key_b), na=False) &
                pool_norm_c.str.contains(re.escape(key_c), na=False)
            )

            hit_df = self.pool.loc[cond, [pool_col_G, pool_col_D, pool_col_H, pool_col_A]]

            if hit_df.empty:
                row_hits[idx] = []
                miss_count += 1
                continue

            hit_count += 1

            if len(hit_df) > 1:
                multi_count += 1

            groups: List[List[Any]] = []
            row_seen_product_keys = set()

            for pool_idx, prow in hit_df.iterrows():
                brand = prow[pool_col_G]
                order_no = prow[pool_col_D]
                min_pack = prow[pool_col_H] if pd.notna(prow[pool_col_H]) else ""
                material_code = prow[pool_col_A]

                try:
                    pool_excel_row = int(pool_idx) + 2
                except Exception:
                    pool_excel_row = self.pool.index.get_loc(pool_idx) + 2

                bom_excel_row = idx + 2

                unique_key = "|".join([
                    self._norm(brand),
                    self._norm(order_no),
                    self._norm(min_pack),
                    self._norm(material_code),
                ])

                if unique_key not in unique_products:
                    unique_products[unique_key] = {
                        "品牌": brand,
                        "订购号": order_no,
                        "最小包装": min_pack,
                        "物料编码": material_code,
                        "命中次数": 0,
                        "BOM行号": [],
                        "汇总表行号": [],
                        "重复匹配汇总表行号": [],
                        "已写入采购表": False,
                    }

                item = unique_products[unique_key]

                if pool_excel_row not in item["汇总表行号"]:
                    item["汇总表行号"].append(pool_excel_row)

                if len(item["汇总表行号"]) > 1:
                    item["重复匹配汇总表行号"] = item["汇总表行号"][:]

                if unique_key not in row_seen_product_keys:
                    row_seen_product_keys.add(unique_key)

                    if bom_excel_row not in item["BOM行号"]:
                        item["BOM行号"].append(bom_excel_row)
                        item["命中次数"] = len(item["BOM行号"])

                    if not item["已写入采购表"]:
                        groups.append([brand, order_no, min_pack, material_code])
                        item["已写入采购表"] = True
                    else:
                        repeat_skip_count += 1
                else:
                    repeat_skip_count += 1

            row_hits[idx] = groups

            if (idx + 1) % 200 == 0:
                self.logger.info(f"匹配进度：{idx + 1}/{len(self.bom)}")

        self.row_hits = row_hits
        self.unique_products = unique_products

        duplicate_pool_product_count = sum(
            1 for item in unique_products.values()
            if len(item.get("汇总表行号", [])) > 1
        )

        self.stats = {
            "total_rows": len(self.bom),
            "hit_count": hit_count,
            "unique_product_count": len(unique_products),
            "repeat_skip_count": repeat_skip_count,
            "multi_count": multi_count,
            "miss_count": miss_count,
            "duplicate_pool_product_count": duplicate_pool_product_count,
            "max_group_num": max((len(v) for v in row_hits.values()), default=0),
            "output_sheet": self.cfg["output_sheet"],
            "summary_sheet": self.cfg["summary_sheet"],
            "bom_file": str(self.bom_file),
        }

    def _write_output_sheet(self) -> None:
        wb = load_workbook(self.bom_file)
        src_ws = wb[wb.sheetnames[0]]
        output_name = self.cfg["output_sheet"]

        if output_name in wb.sheetnames:
            del wb[output_name]

        new_ws = wb.copy_worksheet(src_ws)
        new_ws.title = output_name

        start_col = src_ws.max_column + 1
        headers = self.cfg["group_headers"]
        group_width = len(headers)
        max_group_num = self.stats.get("max_group_num", 0)

        for g in range(max_group_num):
            fill_color = self.cfg["header_fill_even"] if g % 2 == 0 else self.cfg["header_fill_odd"]
            fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

            for k, h in enumerate(headers):
                col = start_col + g * group_width + k
                cell = new_ws.cell(row=1, column=col, value=h)
                cell.fill = fill

        for idx, groups in self.row_hits.items():
            excel_row = idx + 2
            for g, vals in enumerate(groups):
                base_col = start_col + g * group_width
                for k, v in enumerate(vals):
                    new_ws.cell(
                        row=excel_row,
                        column=base_col + k,
                        value=v if pd.notna(v) else None
                    )

        if self.cfg.get("apply_all_border", True):
            thin = Side(style="thin", color="000000")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            for row in new_ws.iter_rows(
                min_row=1,
                max_row=new_ws.max_row,
                min_col=1,
                max_col=new_ws.max_column
            ):
                for cell in row:
                    cell.border = border

        self._write_summary_sheet(wb)

        wb.save(self.bom_file)
        self.logger.success(f"写入完成：{self.bom_file}")

    def _write_summary_sheet(self, wb) -> None:
        sheet_name = self.cfg["summary_sheet"]

        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

        ws = wb.create_sheet(sheet_name)

        headers = [
            "品牌",
            "订购号",
            "最小包装",
            "物料编码",
            "BOM命中次数",
            "BOM行号",
            "汇总表匹配行号",
            "是否汇总表重复匹配",
            "重复匹配汇总表行号",
        ]

        for col, h in enumerate(headers, start=1):
            ws.cell(row=1, column=col, value=h)

        for r, item in enumerate(self.unique_products.values(), start=2):
            bom_rows = item.get("BOM行号", [])
            pool_rows = item.get("汇总表行号", [])
            duplicate_pool_rows = item.get("重复匹配汇总表行号", [])

            is_duplicate_pool_match = "是" if len(pool_rows) > 1 else "否"

            ws.cell(row=r, column=1, value=item.get("品牌"))
            ws.cell(row=r, column=2, value=item.get("订购号"))
            ws.cell(row=r, column=3, value=item.get("最小包装"))
            ws.cell(row=r, column=4, value=item.get("物料编码"))
            ws.cell(row=r, column=5, value=item.get("命中次数"))
            ws.cell(row=r, column=6, value=",".join(map(str, bom_rows)))
            ws.cell(row=r, column=7, value=",".join(map(str, pool_rows)))
            ws.cell(row=r, column=8, value=is_duplicate_pool_match)
            ws.cell(row=r, column=9, value=",".join(map(str, duplicate_pool_rows)))

        if self.cfg.get("apply_all_border", True):
            thin = Side(style="thin", color="000000")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            for row in ws.iter_rows(
                min_row=1,
                max_row=ws.max_row,
                min_col=1,
                max_col=ws.max_column
            ):
                for cell in row:
                    cell.border = border

    @staticmethod
    def _norm(s: Any) -> str:
        if s is None or (isinstance(s, float) and pd.isna(s)):
            return ""

        t = str(s).strip().upper()
        t = re.sub(r"\s+", "", t)

        t = t.replace("（", "(").replace("）", ")")
        t = t.replace("／", "/").replace("\\", "/")
        t = t.replace("－", "-").replace("—", "-")

        t = re.sub(
            r"(电阻)(\d+(?:\.\d+)?)(?:R|Ω)?",
            lambda m: f"{m.group(1)}{m.group(2)}",
            t
        )

        return t


if __name__ == "__main__":
    bom_file = Path(r"VX_MD_POWER_V1.0_BOM.xlsx")
    pool_file = Path(r"电子料汇总表-V1.0.xlsx")

    logger = AppLogger()

    matcher = BOMMatcher(
        bom_file=bom_file,
        pool_file=pool_file,
        config={
            "pool_target_sheet": "26年正式购买",
            "output_sheet": "采购表",
            "summary_sheet": "匹配统计",
            "apply_all_border": True,
        },
        logger=logger
    )

    stats = matcher.run()
    print(stats)